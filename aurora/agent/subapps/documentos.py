"""Aurora Documentos: crea archivos de Word (.docx) sin abrir Word."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .comun import ErrorSubApp, huella, preparar_destino

TIPOS = {"encabezado", "parrafo", "lista", "lista_numerada", "tabla", "salto_pagina"}


def crear_documento(
    destino: Path,
    titulo: str,
    bloques: list[dict],
    subtitulo: str | None = None,
    sobrescribir: bool = False,
) -> dict:
    destino = preparar_destino(destino, ".docx", sobrescribir)
    if not bloques:
        raise ErrorSubApp("El documento no tiene contenido.")

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = titulo
    doc.core_properties.author = "Aurora"

    doc.add_heading(titulo, level=0)
    if subtitulo:
        doc.add_paragraph(subtitulo, style="Subtitle")

    for i, bloque in enumerate(bloques, start=1):
        tipo = bloque.get("tipo")
        if tipo not in TIPOS:
            raise ErrorSubApp(f"Bloque {i}: tipo '{tipo}' no válido. Usa uno de {sorted(TIPOS)}.")
        if tipo == "encabezado":
            nivel = min(max(int(bloque.get("nivel", 1)), 1), 4)
            doc.add_heading(bloque.get("texto", ""), level=nivel)
        elif tipo == "parrafo":
            _parrafo(doc, bloque.get("texto", ""))
        elif tipo in ("lista", "lista_numerada"):
            estilo = "List Bullet" if tipo == "lista" else "List Number"
            for item in bloque.get("items", []):
                _parrafo(doc, str(item), estilo)
        elif tipo == "tabla":
            _tabla(doc, bloque, i)
        elif tipo == "salto_pagina":
            doc.add_page_break()

    doc.save(destino)
    Document(destino)  # verifica que el archivo se puede volver a abrir
    return {"ruta": str(destino), "verificado": True, **huella(destino)}


def _parrafo(doc, texto: str, estilo: str | None = None) -> None:
    """Agrega un párrafo; el texto entre **dobles asteriscos** va en negritas."""
    parrafo = doc.add_paragraph(style=estilo)
    for j, trozo in enumerate(texto.split("**")):
        if trozo:
            parrafo.add_run(trozo).bold = j % 2 == 1


def _tabla(doc, bloque: dict, i: int) -> None:
    encabezados = [str(c) for c in bloque.get("encabezados", [])]
    filas = bloque.get("filas", [])
    if not encabezados:
        raise ErrorSubApp(f"Bloque {i}: la tabla necesita encabezados.")
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        celda.text = texto
        for run in celda.paragraphs[0].runs:
            run.bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = "" if valor is None else str(valor)
    if bloque.get("nota"):
        doc.add_paragraph(bloque["nota"], style="Caption")
