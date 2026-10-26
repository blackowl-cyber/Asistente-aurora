"""Aurora Hojas: crea hojas de cálculo de Excel (.xlsx) sin abrir Excel."""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comun import ErrorSubApp, huella, preparar_destino

RELLENO_ENCABEZADO = PatternFill("solid", fgColor="1F4E78")


def crear_hoja(destino: Path, hojas: list[dict], sobrescribir: bool = False) -> dict:
    destino = preparar_destino(destino, ".xlsx", sobrescribir)
    if not hojas:
        raise ErrorSubApp("El libro necesita al menos una hoja.")

    libro = Workbook()
    libro.remove(libro.active)
    for i, datos in enumerate(hojas, start=1):
        _hoja(libro, datos, i)

    libro.save(destino)
    load_workbook(destino)  # verifica que el archivo se puede volver a abrir
    return {
        "ruta": str(destino),
        "verificado": True,
        "hojas": [h.get("nombre") or f"Hoja{i}" for i, h in enumerate(hojas, start=1)],
        **huella(destino),
    }


def _hoja(libro: Workbook, datos: dict, i: int) -> None:
    nombre = (datos.get("nombre") or f"Hoja{i}")[:31]
    encabezados = datos.get("encabezados") or []
    filas = datos.get("filas") or []
    if not encabezados:
        raise ErrorSubApp(f"Hoja '{nombre}': faltan los encabezados.")
    for n, fila in enumerate(filas, start=1):
        if len(fila) > len(encabezados):
            raise ErrorSubApp(f"Hoja '{nombre}', fila {n}: tiene más columnas que encabezados.")

    hoja = libro.create_sheet(nombre)
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = RELLENO_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append(fila)  # un texto que empieza con "=" se guarda como fórmula

    for columna, formato in (datos.get("formatos") or {}).items():
        indice = encabezados.index(columna) + 1 if columna in encabezados else None
        if indice is None:
            raise ErrorSubApp(f"Hoja '{nombre}': la columna '{columna}' no existe.")
        for (celda,) in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
            celda.number_format = formato

    for indice, titulo in enumerate(encabezados, start=1):
        valores = [str(titulo)] + [str(f[indice - 1]) for f in filas if indice <= len(f)]
        ancho = min(max(len(v) for v in valores) + 2, 60)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = hoja.dimensions

    if datos.get("grafica") and filas:
        _grafica(hoja, datos["grafica"], encabezados, len(filas))


def _grafica(hoja, config: dict, encabezados: list, n_filas: int) -> None:
    grafica = LineChart() if config.get("tipo") == "lineas" else BarChart()
    grafica.title = config.get("titulo", "")
    grafica.height, grafica.width = 9, 18

    def indice(columna: str) -> int:
        if columna not in encabezados:
            raise ErrorSubApp(f"Gráfica: la columna '{columna}' no existe.")
        return encabezados.index(columna) + 1

    for columna in config.get("valores", []):
        col = indice(columna)
        datos = Reference(hoja, min_col=col, min_row=1, max_row=n_filas + 1)
        grafica.add_data(datos, titles_from_data=True)
    col_cat = indice(config["categorias"])
    grafica.set_categories(Reference(hoja, min_col=col_cat, min_row=2, max_row=n_filas + 1))
    hoja.add_chart(grafica, f"{get_column_letter(len(encabezados) + 2)}2")
