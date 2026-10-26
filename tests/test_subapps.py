import pytest
from docx import Document
from openpyxl import load_workbook

from aurora.agent.herramientas import Herramientas
from aurora.agent.sandbox import RutaNoPermitida, Sandbox
from aurora.agent.subapps.comun import ErrorSubApp


@pytest.fixture
def h(tmp_path):
    return Herramientas(Sandbox(tmp_path / "Aurora"))


def test_sandbox_rechaza_rutas_fuera(tmp_path):
    sb = Sandbox(tmp_path / "Aurora")
    assert sb.resolver("Proyectos/a.docx") == sb.raiz / "Proyectos" / "a.docx"
    for mala in ["../fuera.txt", "Proyectos/../../x", str(tmp_path / "otro.txt"), ""]:
        with pytest.raises(RutaNoPermitida):
            sb.resolver(mala)


def test_crear_documento(h):
    r = h.ejecutar("crear_documento", {
        "ruta": "Proyectos/Prueba/informe",
        "titulo": "Informe de materiales",
        "bloques": [
            {"tipo": "encabezado", "texto": "Resumen", "nivel": 1},
            {"tipo": "parrafo", "texto": "El proveedor **más barato** es A."},
            {"tipo": "lista", "items": ["uno", "dos"]},
            {"tipo": "tabla", "encabezados": ["Proveedor", "Precio"], "filas": [["A", 10], ["B", 12.5]], "nota": "Fuente: sitios oficiales"},
        ],
    })
    assert r["ruta"] == "Proyectos/Prueba/informe.docx"
    assert r["verificado"] and r["bytes"] > 0
    doc = Document(h.sb.raiz / r["ruta"])
    textos = [p.text for p in doc.paragraphs]
    assert "Informe de materiales" in textos and "Resumen" in textos
    assert doc.tables[0].cell(2, 1).text == "12.5"


def test_documento_no_sobrescribe(h):
    args = {"ruta": "a.docx", "titulo": "A", "bloques": [{"tipo": "parrafo", "texto": "x"}]}
    h.ejecutar("crear_documento", args)
    with pytest.raises(ErrorSubApp):
        h.ejecutar("crear_documento", args)
    assert h.ejecutar("crear_documento", {**args, "sobrescribir": True})["verificado"]


def test_documento_bloque_invalido(h):
    with pytest.raises(ErrorSubApp):
        h.ejecutar("crear_documento", {"ruta": "a.docx", "titulo": "A", "bloques": [{"tipo": "video"}]})


def test_crear_hoja_con_formulas_y_grafica(h):
    r = h.ejecutar("crear_hoja", {
        "ruta": "Proyectos/Prueba/precios.xlsx",
        "hojas": [{
            "nombre": "Comparativa",
            "encabezados": ["Proveedor", "Precio", "Cantidad", "Total"],
            "filas": [["A", 10, 3, "=B2*C2"], ["B", 12.5, 2, "=B3*C3"]],
            "formatos": {"Precio": "$#,##0.00"},
            "grafica": {"tipo": "barras", "titulo": "Totales", "categorias": "Proveedor", "valores": ["Total"]},
        }],
    })
    assert r["verificado"] and r["hojas"] == ["Comparativa"]
    hoja = load_workbook(h.sb.raiz / r["ruta"])["Comparativa"]
    assert hoja["D2"].value == "=B2*C2"
    assert hoja["B2"].number_format == "$#,##0.00"
    assert hoja.freeze_panes == "A2"
    assert len(hoja._charts) == 1


def test_hoja_columna_inexistente(h):
    with pytest.raises(ErrorSubApp):
        h.ejecutar("crear_hoja", {"ruta": "x.xlsx", "hojas": [{"encabezados": ["A"], "filas": [[1]], "formatos": {"B": "0"}}]})


def test_archivos_listar_mover(h):
    h.ejecutar("crear_carpeta", {"ruta": "Proyectos/Viejo"})
    (h.sb.raiz / "nota.txt").write_text("hola")
    r = h.ejecutar("mover_archivo", {"origen": "nota.txt", "destino": "Proyectos/Viejo"})
    assert r["destino"] == "Proyectos/Viejo/nota.txt" and r["verificado"]
    rutas = [e["ruta"] for e in h.ejecutar("listar_archivos", {"recursivo": True})["entradas"]]
    assert "Proyectos/Viejo/nota.txt" in rutas
    (h.sb.raiz / "nota.txt").write_text("otra")
    with pytest.raises(ErrorSubApp):
        h.ejecutar("mover_archivo", {"origen": "nota.txt", "destino": "Proyectos/Viejo"})


def test_leer_para_envio(h):
    (h.sb.raiz / "a.txt").write_bytes(b"abc")
    assert h.ejecutar("leer_para_envio", {"ruta": "a.txt"}) == {"nombre": "a.txt", "contenido_b64": "YWJj"}


def test_abrir_rechaza_esquemas_raros(h):
    with pytest.raises(ErrorSubApp):
        h.ejecutar("abrir", {"objetivo": "file.exe"})
